from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from partners.models import Partner,PartnerCategory
class Command(BaseCommand):
    help='협력업체 엑셀 파일을 DB에 등록/갱신합니다.'
    def add_arguments(self,p):
        p.add_argument('xlsx'); p.add_argument('--category',default='grooming'); p.add_argument('--name',default='미용센터')
    def handle(self,*args,**o):
        cat,_=PartnerCategory.objects.get_or_create(code=o['category'],defaults={'name':o['name']})
        wb=load_workbook(o['xlsx'],data_only=True); created=updated=0
        for ws in wb.worksheets:
            city=ws.title.rstrip('0123456789')
            for row in ws.iter_rows(min_row=4,values_only=True):
                if len(row)<3 or not row[1] or not row[2]: continue
                _,new=Partner.objects.update_or_create(category=cat,name=str(row[1]).strip(),address=str(row[2]).strip(),defaults={'phone':str(row[3]).strip() if len(row)>3 and row[3] else '','city':city,'source':o['xlsx'],'is_active':True})
                created+=new; updated+=not new
        self.stdout.write(self.style.SUCCESS(f'완료: 신규 {created}, 갱신 {updated}'))
