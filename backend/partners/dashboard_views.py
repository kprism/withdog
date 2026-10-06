from django.http import Http404
from django.views.decorators.http import require_POST
from django.contrib.auth.decorators import login_required
from pathlib import Path
from tempfile import NamedTemporaryFile
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from django.db.models import Count,Q
from django.shortcuts import redirect,render,get_object_or_404
from django.utils.dateparse import parse_date,parse_datetime
from openpyxl import load_workbook
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
import json
from .models import Partner,PartnerCategory,PartnerBenefit,ContentItem,MemberProfile,DogRegistration,MemberBenefit,Inquiry, SiteSetting, DogBreed

LABELS={'notice':'공지사항','community':'자유게시판','breed':'전 세계 견종','intro':'인트로 갤러리','popup':'팝업'}

def _kakao_geocode_address(address):
    """
    카카오 Local REST API를 이용하여
    도로명/지번 주소를 위도(latitude), 경도(longitude)로 변환한다.

    실패 시 예외를 외부로 던지지 않고 (None, None)을 반환한다.
    """
    address = str(address or '').strip()

    if not address:
        return None, None

    setting = SiteSetting.get_solo()
    rest_key = (setting.kakao_rest_api_key or '').strip()

    if not rest_key:
        return None, None

    query = urlencode({
        'query': address
    })

    url = (
        'https://dapi.kakao.com/v2/local/search/address.json?'
        + query
    )

    req = Request(
        url,
        headers={
            'Authorization': 'KakaoAK ' + rest_key,
            'User-Agent': 'ThePetKorea/1.0',
        }
    )

    try:
        with urlopen(req, timeout=8) as response:
            payload = json.loads(
                response.read().decode('utf-8')
            )

        documents = payload.get('documents') or []

        if not documents:
            return None, None

        first = documents[0]

        # Kakao:
        # x = longitude
        # y = latitude
        latitude = first.get('y')
        longitude = first.get('x')

        if not latitude or not longitude:
            return None, None

        return latitude, longitude

    except (HTTPError, URLError, TimeoutError, ValueError, OSError):
        return None, None



def _kakao_find_place(name, address, latitude=None, longitude=None):
    """
    카카오 Local 키워드 검색으로 실제 카카오 장소를 찾는다.

    업체명 우선 검색 후 주소/좌표를 이용하여
    가장 적절한 결과의 place id / place_url을 반환한다.
    """
    name = str(name or '').strip()
    address = str(address or '').strip()

    if not name:
        return '', ''

    setting = SiteSetting.get_solo()
    rest_key = (setting.kakao_rest_api_key or '').strip()

    if not rest_key:
        return '', ''

    params = {
        'query': name,
        'size': 15,
    }

    if longitude is not None and latitude is not None:
        params['x'] = str(longitude)
        params['y'] = str(latitude)
        params['radius'] = 20000
        params['sort'] = 'distance'

    url = (
        'https://dapi.kakao.com/v2/local/search/keyword.json?'
        + urlencode(params)
    )

    req = Request(
        url,
        headers={
            'Authorization': 'KakaoAK ' + rest_key,
            'User-Agent': 'ThePetKorea/1.0',
        }
    )

    try:
        with urlopen(req, timeout=8) as response:
            payload = json.loads(
                response.read().decode('utf-8')
            )

        documents = payload.get('documents') or []

        if not documents:
            return '', ''

        def normalize(value):
            return ''.join(
                str(value or '').lower().split()
            )

        target_name = normalize(name)
        target_address = normalize(address)

        best = None
        best_score = -1

        for doc in documents:
            place_name = normalize(doc.get('place_name'))
            road_address = normalize(doc.get('road_address_name'))
            jibun_address = normalize(doc.get('address_name'))

            score = 0

            if place_name == target_name:
                score += 100
            elif target_name and target_name in place_name:
                score += 60
            elif place_name and place_name in target_name:
                score += 50

            if target_address:
                if road_address == target_address:
                    score += 100
                elif jibun_address == target_address:
                    score += 100
                elif (
                    road_address
                    and (
                        road_address in target_address
                        or target_address in road_address
                    )
                ):
                    score += 50
                elif (
                    jibun_address
                    and (
                        jibun_address in target_address
                        or target_address in jibun_address
                    )
                ):
                    score += 50

            try:
                distance = int(doc.get('distance') or 999999)
            except (TypeError, ValueError):
                distance = 999999

            if distance <= 100:
                score += 40
            elif distance <= 300:
                score += 30
            elif distance <= 1000:
                score += 20
            elif distance <= 3000:
                score += 10

            if score > best_score:
                best_score = score
                best = doc

        # 너무 약한 매칭은 저장하지 않는다.
        if not best or best_score < 50:
            return '', ''

        place_id = str(best.get('id') or '').strip()
        place_url = str(best.get('place_url') or '').strip()

        if place_id and not place_url:
            place_url = (
                'https://place.map.kakao.com/' + place_id
            )

        return place_id, place_url

    except (
        HTTPError,
        URLError,
        TimeoutError,
        ValueError,
        OSError,
    ):
        return '', ''



def _kakao_find_place_secondary(
    name,
    address,
    phone='',
    latitude=None,
    longitude=None,
):
    """
    1차 업체명 검색에서 실패한 업체를 위한 보조 검색.

    검색 전략:
    1. 업체명
    2. 업체명 + 주소의 시/군/구
    3. 주소
    4. 전화번호

    이 함수는 후보를 반환할 뿐 자동 저장하지 않는다.
    """

    import re
    import json

    name = str(name or '').strip()
    address = str(address or '').strip()
    phone = str(phone or '').strip()

    setting = SiteSetting.get_solo()
    rest_key = (setting.kakao_rest_api_key or '').strip()

    if not rest_key:
        return []

    # 주소에서 지역명 추출
    region_parts = re.findall(
        r'[가-힣]+(?:시|군|구|읍|면|동)',
        address
    )

    region = ' '.join(region_parts[:2])

    queries = []

    if name:
        queries.append(name)

    if name and region:
        queries.append(
            f'{name} {region}'
        )

    if address:
        # 건물 상세호수 등은 오히려 검색을 방해할 수 있으므로
        # 전체 주소와 앞부분 주소를 둘 다 시도
        queries.append(address)

        address_tokens = address.split()

        if len(address_tokens) >= 4:
            queries.append(
                ' '.join(address_tokens[:4])
            )

    if phone:
        queries.append(phone)

    # 중복 제거
    unique_queries = []

    for q in queries:
        q = q.strip()

        if q and q not in unique_queries:
            unique_queries.append(q)

    candidates = {}

    def normalize(v):
        return re.sub(
            r'[^0-9a-z가-힣]',
            '',
            str(v or '').lower()
        )

    target_name = normalize(name)
    target_address = normalize(address)

    target_phone = re.sub(
        r'[^0-9]',
        '',
        phone
    )

    for query in unique_queries:

        params = {
            'query': query,
            'size': 15,
        }

        # 좌표가 있으면 검색 중심으로 사용
        if (
            latitude is not None
            and longitude is not None
        ):
            params['x'] = str(longitude)
            params['y'] = str(latitude)
            params['radius'] = 20000
            params['sort'] = 'distance'

        url = (
            'https://dapi.kakao.com/v2/local/search/keyword.json?'
            + urlencode(params)
        )

        req = Request(
            url,
            headers={
                'Authorization':
                    'KakaoAK ' + rest_key,
                'User-Agent':
                    'ThePetKorea/1.0',
            }
        )

        try:

            with urlopen(
                req,
                timeout=8
            ) as response:

                payload = json.loads(
                    response
                    .read()
                    .decode('utf-8')
                )

        except Exception:
            continue

        for doc in payload.get(
            'documents',
            []
        ):

            place_id = str(
                doc.get('id') or ''
            ).strip()

            if not place_id:
                continue

            place_name = str(
                doc.get('place_name') or ''
            )

            road_address = str(
                doc.get('road_address_name') or ''
            )

            jibun_address = str(
                doc.get('address_name') or ''
            )

            place_phone = str(
                doc.get('phone') or ''
            )

            place_url = str(
                doc.get('place_url') or ''
            )

            if not place_url:
                place_url = (
                    'https://place.map.kakao.com/'
                    + place_id
                )

            n_name = normalize(place_name)
            n_road = normalize(road_address)
            n_jibun = normalize(jibun_address)

            n_phone = re.sub(
                r'[^0-9]',
                '',
                place_phone
            )

            score = 0
            reasons = []

            # 업체명
            if (
                target_name
                and n_name == target_name
            ):
                score += 100
                reasons.append('상호일치')

            elif (
                target_name
                and (
                    target_name in n_name
                    or n_name in target_name
                )
            ):
                score += 55
                reasons.append('상호유사')

            # 주소
            if target_address:

                if (
                    n_road
                    and n_road == target_address
                ):
                    score += 100
                    reasons.append('도로명주소일치')

                elif (
                    n_jibun
                    and n_jibun == target_address
                ):
                    score += 100
                    reasons.append('지번주소일치')

                elif (
                    n_road
                    and (
                        n_road in target_address
                        or target_address in n_road
                    )
                ):
                    score += 50
                    reasons.append('도로명주소유사')

                elif (
                    n_jibun
                    and (
                        n_jibun in target_address
                        or target_address in n_jibun
                    )
                ):
                    score += 50
                    reasons.append('지번주소유사')

            # 전화번호
            if (
                target_phone
                and n_phone
                and target_phone == n_phone
            ):
                score += 150
                reasons.append('전화번호일치')

            # 거리
            try:
                distance = int(
                    doc.get('distance') or 999999
                )
            except Exception:
                distance = 999999

            if distance <= 50:
                score += 80
                reasons.append('50m이내')

            elif distance <= 100:
                score += 60
                reasons.append('100m이내')

            elif distance <= 300:
                score += 40
                reasons.append('300m이내')

            elif distance <= 1000:
                score += 20
                reasons.append('1km이내')

            item = {
                'id': place_id,
                'name': place_name,
                'address': (
                    road_address
                    or jibun_address
                ),
                'phone': place_phone,
                'url': place_url,
                'distance': distance,
                'score': score,
                'reasons': reasons,
                'query': query,
            }

            previous = candidates.get(
                place_id
            )

            if (
                previous is None
                or score > previous['score']
            ):
                candidates[place_id] = item

    result = list(
        candidates.values()
    )

    result.sort(
        key=lambda x: (
            -x['score'],
            x['distance'],
        )
    )

    return result[:5]


def _import_partner_workbook(uploaded_file, category):
    created = 0
    updated = 0
    skipped = 0
    geocoded = 0
    geocode_failed = 0

    with NamedTemporaryFile(
        suffix=Path(uploaded_file.name).suffix or '.xlsx'
    ) as tmp:

        for chunk in uploaded_file.chunks():
            tmp.write(chunk)

        tmp.flush()

        wb = load_workbook(
            tmp.name,
            data_only=True,
            read_only=True
        )

        for ws in wb.worksheets:

            # 시트명 예:
            # 창원시1 -> 창원시
            # 김해시2 -> 김해시
            city = ws.title.rstrip('0123456789').strip()

            for row in ws.iter_rows(
                min_row=4,
                values_only=True
            ):
                if (
                    len(row) < 3
                    or not row[1]
                    or not row[2]
                ):
                    skipped += 1
                    continue

                name = str(row[1]).strip()
                address = str(row[2]).strip()

                phone = (
                    str(row[3]).strip()
                    if len(row) > 3 and row[3]
                    else ''
                )

                partner, new = Partner.objects.update_or_create(
                    category=category,
                    name=name,
                    address=address,
                    defaults={
                        'phone': phone,
                        'city': city,
                        'source': uploaded_file.name,
                        'is_active': True,
                    }
                )

                if new:
                    created += 1
                else:
                    updated += 1

                # ------------------------------------------------
                # 좌표가 없는 경우에만 카카오 API 호출
                # ------------------------------------------------
                if (
                    partner.latitude is None
                    or partner.longitude is None
                ):
                    latitude, longitude = _kakao_geocode_address(
                        address
                    )

                    if latitude and longitude:
                        partner.latitude = latitude
                        partner.longitude = longitude

                        partner.save(
                            update_fields=[
                                'latitude',
                                'longitude',
                                'updated_at',
                            ]
                        )

                        geocoded += 1

                    else:
                        geocode_failed += 1

                # KAKAO PLACE AUTO MATCH
                if not partner.kakao_place_url:
                    place_id, place_url = _kakao_find_place(
                        partner.name,
                        partner.address,
                        partner.latitude,
                        partner.longitude,
                    )

                    if place_url:
                        partner.kakao_place_id = place_id
                        partner.kakao_place_url = place_url
                        partner.save(
                            update_fields=[
                                'kakao_place_id',
                                'kakao_place_url',
                                'updated_at',
                            ]
                        )

    return (
        created,
        updated,
        skipped,
        geocoded,
        geocode_failed,
    )

@staff_member_required
def dashboard_home(request):
 total=Partner.objects.count(); active=Partner.objects.filter(is_active=True).count(); affiliated=Partner.objects.filter(is_affiliated=True).count(); no_coordinates=Partner.objects.filter(Q(latitude__isnull=True)|Q(longitude__isnull=True)).count(); categories=PartnerCategory.objects.annotate(partner_count=Count('partners')).order_by('sort_order','name'); recent=Partner.objects.select_related('category').order_by('-updated_at')[:8]
 return render(request,'dashboard/home.html',locals())

@staff_member_required
def partner_list(request):
    return redirect('operator_dashboard:partner_page_editor')

@staff_member_required
def partner_edit(request,pk=None):
 p=get_object_or_404(Partner,pk=pk) if pk else None; categories=PartnerCategory.objects.order_by('sort_order','name')
 if request.method=='POST':
  obj=p or Partner(); obj.name=request.POST.get('name','').strip(); obj.address=request.POST.get('address','').strip(); obj.phone=request.POST.get('phone','').strip(); obj.city=request.POST.get('city','').strip(); obj.category=get_object_or_404(PartnerCategory,pk=request.POST.get('category')); obj.latitude=request.POST.get('latitude') or None; obj.longitude=request.POST.get('longitude') or None; obj.is_affiliated='is_affiliated' in request.POST; obj.is_active='is_active' in request.POST; obj.save(); messages.success(request,'업체 정보가 저장되었습니다.'); return redirect('operator_dashboard:partners')
 return render(request,'dashboard/partner_form.html',locals())

@staff_member_required
def partner_delete(request,pk):
 if request.method=='POST': get_object_or_404(Partner,pk=pk).delete(); messages.success(request,'업체가 삭제되었습니다.')
 return redirect('operator_dashboard:partners')

@staff_member_required
def partner_upload(request):
    return redirect('operator_dashboard:partner_page_editor')

@staff_member_required
def content_list(request,kind):
 if kind not in LABELS:return redirect('operator_dashboard:home')
 label=LABELS[kind]; q=request.GET.get('q','').strip(); items=ContentItem.objects.filter(kind=kind)
 if q:items=items.filter(Q(title__icontains=q)|Q(body__icontains=q)|Q(label__icontains=q))
 return render(request,'dashboard/content_list.html',locals())
@staff_member_required
def content_edit(request,kind,pk=None):
 if kind not in LABELS:return redirect('operator_dashboard:home')
 label=LABELS[kind]; item=get_object_or_404(ContentItem,pk=pk,kind=kind) if pk else None
 if request.method=='POST':
  obj=item or ContentItem(kind=kind); obj.title=request.POST.get('title','').strip(); obj.label=request.POST.get('label','').strip(); obj.body=request.POST.get('body','').strip(); obj.link=request.POST.get('link','').strip(); obj.sort_order=int(request.POST.get('sort_order') or 0); obj.is_published='is_published' in request.POST
  if request.FILES.get('image'):obj.image=request.FILES['image']
  obj.popup_start=parse_datetime(request.POST.get('popup_start','')) if request.POST.get('popup_start') else None; obj.popup_end=parse_datetime(request.POST.get('popup_end','')) if request.POST.get('popup_end') else None; obj.save(); messages.success(request,'저장되었습니다. 홈페이지에 즉시 반영됩니다.'); return redirect('operator_dashboard:content_list',kind=kind)
 popup_start=item.popup_start.strftime('%Y-%m-%dT%H:%M') if item and item.popup_start else ''; popup_end=item.popup_end.strftime('%Y-%m-%dT%H:%M') if item and item.popup_end else ''
 return render(request,'dashboard/content_form.html',locals())
@staff_member_required
def content_delete(request,kind,pk):
 if request.method=='POST':get_object_or_404(ContentItem,pk=pk,kind=kind).delete();messages.success(request,'삭제되었습니다.')
 return redirect('operator_dashboard:content_list',kind=kind)

def _section_config(section):
 return {
  'dogs':{'title':'반려견 등록 관리','model':DogRegistration,'search':['dog_name','registration_no','breed','owner__email'],'headers':['반려견','등록번호','견종','보호자','상태','등록일'],'fields':[('owner','보호자','user'),('dog_name','반려견 이름','text'),('registration_no','동물등록번호','text'),('breed','견종','text'),('birth_date','생년월일','date'),('gender','성별','text'),('status','처리상태','dog_status')]},
  'benefits':{'title':'회원 혜택 관리','model':MemberBenefit,'search':['title','description'],'headers':['혜택명','내용','시작일','종료일','상태','순서'],'fields':[('title','혜택명','text'),('description','혜택 내용','textarea'),('valid_from','시작일','date'),('valid_to','종료일','date'),('is_active','노출','checkbox'),('sort_order','정렬순서','number')]},
  'partner-benefits':{'title':'제휴 혜택 관리','model':PartnerBenefit,'search':['partner__name','title','description'],'headers':['업체','혜택명','내용','노출'],'fields':[('partner','제휴업체','partner'),('title','혜택명','text'),('description','혜택 내용','textarea'),('active','노출','checkbox')]},
  'inquiries':{'title':'민원·문의 관리','model':Inquiry,'search':['name','phone','email','subject','body'],'headers':['신청인','연락처','이메일','제목','상태','접수일'],'fields':[('name','신청인','text'),('phone','전화번호','text'),('email','이메일','email'),('subject','제목','text'),('body','문의내용','textarea'),('status','처리상태','inquiry_status'),('admin_memo','관리자 메모','textarea')]},
 } .get(section)

def _row_values(section,obj):
 if section=='dogs': return [obj.dog_name,obj.registration_no or '-',obj.breed or '-',obj.owner.email or obj.owner.username,obj.get_status_display(),obj.created_at.strftime('%Y.%m.%d')]
 if section=='benefits': return [obj.title,obj.description[:50],obj.valid_from or '-',obj.valid_to or '-', '노출' if obj.is_active else '숨김',obj.sort_order]
 if section=='partner-benefits': return [obj.partner.name,obj.title,obj.description[:60],'노출' if obj.active else '숨김']
 if section=='inquiries': return [obj.name,obj.phone or '-',obj.email or '-',obj.subject,obj.get_status_display(),obj.created_at.strftime('%Y.%m.%d')]
 return []

@staff_member_required
def manage_section(request,section):
 redirects={'notices':'notice','freeboard':'community','breeds':'breed','content':'intro'}
 if section in redirects:return redirect('operator_dashboard:content_list',kind=redirects[section])
 if section=='members': return member_list(request)
 if section=='admins': return admin_list(request)
 cfg=_section_config(section)
 if not cfg:return redirect('operator_dashboard:home')
 q=request.GET.get('q','').strip(); qs=cfg['model'].objects.all()
 if q:
  query=Q()
  for f in cfg['search']: query|=Q(**{f+'__icontains':q})
  qs=qs.filter(query)
 rows=[{'id':o.pk,'values':_row_values(section,o)} for o in qs[:200]]
 return render(request,'dashboard/manage_section.html',{'section':section,'section_title':cfg['title'],'headers':cfg['headers'],'rows':rows,'q':q})

@staff_member_required
def section_edit(request,section,pk=None):
 cfg=_section_config(section)
 if not cfg:return redirect('operator_dashboard:home')
 obj=get_object_or_404(cfg['model'],pk=pk) if pk else None
 if request.method=='POST':
  x=obj or cfg['model']()
  for name,label,typ in cfg['fields']:
   v=request.POST.get(name,'')
   if typ=='checkbox': v=name in request.POST
   elif typ=='number': v=int(v or 0)
   elif typ=='date': v=parse_date(v) if v else None
   elif typ=='user': v=get_object_or_404(get_user_model(),pk=v)
   elif typ=='partner': v=get_object_or_404(Partner,pk=v)
   setattr(x,name,v)
  x.save(); messages.success(request,'저장되었습니다.'); return redirect('operator_dashboard:manage_section',section=section)
 users=get_user_model().objects.filter(is_staff=False,is_active=True).order_by('email'); partners=Partner.objects.filter(is_active=True).order_by('name')
 return render(request,'dashboard/section_form.html',{'section':section,'title':cfg['title'],'fields':cfg['fields'],'obj':obj,'users':users,'partners':partners})

@staff_member_required
def section_delete(request,section,pk):
 cfg=_section_config(section)
 if cfg and request.method=='POST': get_object_or_404(cfg['model'],pk=pk).delete(); messages.success(request,'삭제되었습니다.')
 return redirect('operator_dashboard:manage_section',section=section)

@staff_member_required
def member_list(request):
 User=get_user_model(); q=request.GET.get('q','').strip(); status=request.GET.get('status','active').strip(); qs=User.objects.filter(is_staff=False).select_related('member_profile')
 if status == 'withdrawn':
  qs=qs.filter(Q(is_active=False)|Q(member_profile__membership_status='withdrawn'))
 elif status == 'all':
  pass
 else:
  status='active'; qs=qs.filter(is_active=True).exclude(member_profile__membership_status='withdrawn')
 if q: qs=qs.filter(Q(email__icontains=q)|Q(username__icontains=q)|Q(member_profile__name__icontains=q)|Q(member_profile__phone__icontains=q)|Q(member_profile__region__icontains=q))
 rows=[]
 for u in qs.order_by('-date_joined')[:300]:
  p=getattr(u,'member_profile',None); rows.append({'id':u.pk,'name':p.name if p else u.get_full_name() or '-','birth':p.birth_date if p else '-','gender':p.get_gender_display() if p and p.gender else '-','phone':p.phone if p else '-','email':u.email or u.username,'region':p.region if p else '-','address':p.address_detail if p else '-','grade':p.get_membership_status_display() if p else '일반회원','active':u.is_active,'joined':u.date_joined})
 return render(request,'dashboard/member_list.html',{'rows':rows,'q':q})

@staff_member_required
def member_edit(request,pk=None):
 User=get_user_model(); user=get_object_or_404(User,pk=pk,is_staff=False) if pk else None; profile=getattr(user,'member_profile',None) if user else None
 if request.method=='POST':
  email=request.POST.get('email','').strip(); user=user or User(username=email,email=email); user.username=email; user.email=email; user.is_active='is_active' in request.POST
  if request.POST.get('password'): user.set_password(request.POST['password'])
  elif not user.pk: user.set_unusable_password()
  user.save(); p=profile or MemberProfile(user=user); p.name=request.POST.get('name','').strip(); p.birth_date=parse_date(request.POST.get('birth_date','')) if request.POST.get('birth_date') else None; p.gender=request.POST.get('gender',''); p.phone=request.POST.get('phone','').strip(); p.region=request.POST.get('region','').strip(); p.address_detail=request.POST.get('address_detail','').strip(); p.privacy_agreed='privacy_agreed' in request.POST; p.regular_member_requested='regular_member_requested' in request.POST; p.membership_status=request.POST.get('membership_status','general'); p.save(); messages.success(request,'회원 정보가 저장되었습니다.'); return redirect('operator_dashboard:manage_section',section='members')
 return render(request,'dashboard/member_form.html',locals())

@staff_member_required
def member_withdraw(request,pk):
 if request.method=='POST':
  u=get_object_or_404(get_user_model(),pk=pk,is_staff=False); u.is_active=False; u.save(update_fields=['is_active']); p=getattr(u,'member_profile',None)
  if p: p.membership_status='withdrawn'; p.save(update_fields=['membership_status'])
  messages.success(request,'회원이 탈퇴(비활성) 처리되었습니다. 데이터는 이력 보존을 위해 유지됩니다.')
 return redirect('operator_dashboard:manage_section',section='members')

@staff_member_required
def admin_list(request):
 User=get_user_model(); q=request.GET.get('q','').strip(); qs=User.objects.filter(is_staff=True)
 if q: qs=qs.filter(Q(username__icontains=q)|Q(email__icontains=q)|Q(first_name__icontains=q))
 rows=qs.order_by('-is_superuser','username'); return render(request,'dashboard/admin_list.html',locals())
@staff_member_required
def admin_edit(request,pk=None):
 User=get_user_model(); obj=get_object_or_404(User,pk=pk,is_staff=True) if pk else None
 if request.method=='POST':
  u=obj or User(); u.username=request.POST.get('username','').strip(); u.email=request.POST.get('email','').strip(); u.first_name=request.POST.get('name','').strip(); u.is_staff=True; u.is_active='is_active' in request.POST; u.is_superuser='is_superuser' in request.POST
  if request.POST.get('password'): u.set_password(request.POST['password'])
  elif not u.pk: messages.error(request,'신규 관리자는 비밀번호가 필요합니다.'); return render(request,'dashboard/admin_form.html',locals())
  u.save(); messages.success(request,'관리자 계정이 저장되었습니다.'); return redirect('operator_dashboard:manage_section',section='admins')
 return render(request,'dashboard/admin_form.html',locals())



# ============================================================
# PARTNER CATEGORY MANAGEMENT
# ============================================================

@staff_member_required
def partner_category_create(request):
    if request.method != 'POST':
        return redirect('operator_dashboard:partner_page_editor')

    import re
    import uuid

    name = request.POST.get('name', '').strip()

    try:
        sort_order = int(request.POST.get('sort_order') or 0)
    except (TypeError, ValueError):
        sort_order = 0

    if not name:
        messages.error(request, '분류명을 입력해 주세요.')
        return redirect('operator_dashboard:partner_page_editor')

    if PartnerCategory.objects.filter(name=name).exists():
        messages.error(request, '이미 같은 이름의 분류가 있습니다.')
        return redirect('operator_dashboard:partner_page_editor')

    # 영문/숫자가 포함되어 있으면 일부를 활용하고,
    # 한글/이모지 분류명은 안전한 내부코드를 자동 생성한다.
    base = re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')

    if not base:
        base = 'category'

    code = base

    if PartnerCategory.objects.filter(code=code).exists():
        code = f'{base}-{uuid.uuid4().hex[:8]}'

    PartnerCategory.objects.create(
        code=code,
        name=name,
        sort_order=max(0, sort_order)
    )

    messages.success(
        request,
        f'{name} 분류가 추가되었습니다.'
    )

    return redirect('operator_dashboard:partner_page_editor')


@staff_member_required
def partner_category_update(request, pk):
    if request.method != 'POST':
        return redirect('operator_dashboard:partner_page_editor')

    category = get_object_or_404(PartnerCategory, pk=pk)

    name = request.POST.get('name', '').strip()

    try:
        sort_order = int(request.POST.get('sort_order') or 0)
    except (TypeError, ValueError):
        sort_order = 0

    if not name:
        messages.error(request, '분류명을 입력해 주세요.')
        return redirect('operator_dashboard:partner_page_editor')

    if (
        PartnerCategory.objects
        .exclude(pk=category.pk)
        .filter(name=name)
        .exists()
    ):
        messages.error(request, '이미 같은 이름의 분류가 있습니다.')
        return redirect('operator_dashboard:partner_page_editor')

    category.name = name
    category.sort_order = max(0, sort_order)
    category.save(update_fields=['name', 'sort_order'])

    messages.success(
        request,
        f'{category.name} 분류가 수정되었습니다.'
    )

    return redirect('operator_dashboard:partner_page_editor')


@staff_member_required
def partner_category_delete(request, pk):
    if request.method != 'POST':
        return redirect('operator_dashboard:partner_page_editor')

    category = get_object_or_404(PartnerCategory, pk=pk)

    partner_count = category.partners.count()

    if partner_count:
        messages.error(
            request,
            f'{category.name} 분류에는 업체 {partner_count}곳이 등록되어 있어 삭제할 수 없습니다.'
        )
        return redirect('operator_dashboard:partner_page_editor')

    name = category.name
    category.delete()

    messages.success(
        request,
        f'{name} 분류가 삭제되었습니다.'
    )

    return redirect('operator_dashboard:partner_page_editor')



# ============================================================
# RELATED COMPANIES PAGE EDITOR
# ============================================================

@login_required
def partner_page_editor(request):
    site_setting = SiteSetting.get_solo()

    # ---------------------------------------------------------
    # 관련업체 CMS Excel 업로드
    # ---------------------------------------------------------
    if (
        request.method == 'POST'
        and request.POST.get('action') == 'upload_partner_excel'
    ):
        uploaded = request.FILES.get('excel_file')
        category_id = request.POST.get('category', '').strip()

        if not uploaded:
            messages.error(request, '엑셀 파일을 선택해 주세요.')
            return redirect('operator_dashboard:partner_page_editor')

        category = (
            PartnerCategory.objects.filter(pk=category_id).first()
            if category_id
            else None
        )

        if not category:
            messages.error(request, '업로드할 업체 분류를 선택해 주세요.')
            return redirect('operator_dashboard:partner_page_editor')

        try:
            (
                created,
                updated,
                skipped,
                geocoded,
                geocode_failed,
            ) = _import_partner_workbook(
                uploaded,
                category,
            )

        except Exception as exc:
            messages.error(
                request,
                f'엑셀 처리 오류: {exc}'
            )
            return redirect(
                'operator_dashboard:partner_page_editor'
            )

        messages.success(
            request,
            (
                f'{category.name} 업로드 완료 · '
                f'신규 {created} · '
                f'갱신 {updated} · '
                f'좌표생성 {geocoded} · '
                f'좌표실패 {geocode_failed} · '
                f'건너뜀 {skipped}'
            )
        )

        return redirect(
            'operator_dashboard:partner_page_editor'
        )

    if request.method == 'POST' and request.POST.get('action') == 'save_kakao_settings':
        js_key = request.POST.get('kakao_javascript_key', '').strip()
        rest_key = request.POST.get('kakao_rest_api_key', '').strip()

        if js_key:
            site_setting.kakao_javascript_key = js_key
        if rest_key:
            site_setting.kakao_rest_api_key = rest_key

        if request.POST.get('delete_kakao_javascript_key') == '1':
            site_setting.kakao_javascript_key = ''
        if request.POST.get('delete_kakao_rest_api_key') == '1':
            site_setting.kakao_rest_api_key = ''

        site_setting.save()
        messages.success(request, '카카오 지도 API 설정이 저장되었습니다.')
        return redirect('operator_dashboard:partner_page_editor')

    from django.db.models import Count, Q
    from django.utils import timezone

    categories = (
        PartnerCategory.objects
        .annotate(partner_count=Count('partners'))
        .order_by('sort_order', 'name')
    )

    total = Partner.objects.count()
    active = Partner.objects.filter(is_active=True).count()

    today = timezone.localdate()

    affiliated = Partner.objects.filter(
        is_affiliated=True,
        is_active=True
    ).filter(
        Q(ad_start_date__isnull=True) | Q(ad_start_date__lte=today),
        Q(ad_end_date__isnull=True) | Q(ad_end_date__gte=today),
    ).count()

    no_coordinates = Partner.objects.filter(
        Q(latitude__isnull=True) |
        Q(longitude__isnull=True)
    ).count()

    q = request.GET.get('q', '').strip()
    category_id = request.GET.get('category', '').strip()
    city = request.GET.get('city', '').strip()

    partners = Partner.objects.select_related('category').all()

    if q:
        partners = partners.filter(
            Q(name__icontains=q) |
            Q(address__icontains=q) |
            Q(phone__icontains=q)
        )

    if category_id:
        partners = partners.filter(category_id=category_id)

    if city:
        partners = partners.filter(city=city)

    partners = partners.order_by(
        '-is_affiliated',
        '-ad_monthly_fee',
        'city',
        'name'
    )

    cities = (
        Partner.objects
        .exclude(city='')
        .values_list('city', flat=True)
        .distinct()
        .order_by('city')
    )

    context = {
        'site_setting': site_setting,
        'categories': categories,
        'total': total,
        'active': active,
        'affiliated': affiliated,
        'no_coordinates': no_coordinates,
        'partners': partners[:200],
        'cities': cities,
        'q': q,
        'selected_category': category_id,
        'selected_city': city,
    }

    return render(
        request,
        'dashboard/partner_page_editor.html',
        context
    )


@login_required
def partner_ad_update(request, pk):
    if request.method != 'POST':
        return redirect('operator_dashboard:partner_page_editor')

    p = get_object_or_404(Partner, pk=pk)

    from datetime import datetime

    p.is_affiliated = request.POST.get('is_affiliated') == '1'

    try:
        p.ad_monthly_fee = max(
            0,
            int(request.POST.get('ad_monthly_fee') or 0)
        )
    except (TypeError, ValueError):
        p.ad_monthly_fee = 0

    def parse_date(value):
        value = (value or '').strip()
        if not value:
            return None
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except ValueError:
            return None

    p.ad_start_date = parse_date(
        request.POST.get('ad_start_date')
    )
    p.ad_end_date = parse_date(
        request.POST.get('ad_end_date')
    )

    payment = request.POST.get('ad_payment_type', '')
    p.ad_payment_type = (
        payment
        if payment in ('monthly', 'annual')
        else ''
    )

    # 카카오 매장 상세페이지 URL
    #
    # 협력업체일 때 관리자 화면에서 입력한 URL을 저장한다.
    # 일반업체로 변경되어 입력 필드가 POST되지 않는 경우에는
    # 기존 URL을 삭제하지 않고 그대로 보존한다.
    if p.is_affiliated and 'kakao_place_url' in request.POST:
        kakao_place_url = (
            request.POST.get('kakao_place_url') or ''
        ).strip()

        # 카카오맵 장소 상세 URL만 허용
        if kakao_place_url:
            valid_prefixes = (
                'https://place.map.kakao.com/',
                'http://place.map.kakao.com/',
            )

            if not kakao_place_url.startswith(valid_prefixes):
                messages.error(
                    request,
                    '카카오 상세 URL은 place.map.kakao.com 주소만 입력할 수 있습니다.'
                )
                return redirect(
                    'operator_dashboard:partner_page_editor'
                )

        p.kakao_place_url = kakao_place_url

        # URL 마지막 숫자를 place id로 같이 보관
        if kakao_place_url:
            p.kakao_place_id = (
                kakao_place_url
                .rstrip('/')
                .split('/')[-1]
            )

    p.save()

    messages.success(
        request,
        f'{p.name} 협력업체 광고정보가 저장되었습니다.'
    )

    return redirect('operator_dashboard:partner_page_editor')


# ============================================================
# WEBSITE EDITOR - WORLD BREEDS
# ============================================================

@login_required
def breed_page_editor(request):
    """
    홈페이지 편집 > 전 세계 견종

    ContentItem(kind='breed')를 사용한다.
    """

    items = ContentItem.objects.filter(
        kind='breed'
    ).order_by('sort_order', '-created_at')

    dog_breeds = DogBreed.objects.all().order_by(
        'sort_order',
        'name_ko',
        'id',
    )

    return render(
        request,
        'dashboard/breed_page_editor.html',
        {
            'items': items,
            'dog_breeds': dog_breeds,
        }
    )



@login_required
@require_POST
def dogbreed_create(request):
    """
    홈페이지 편집 > 전 세계 견종 > 견종 DB 수동 등록
    """

    name_ko = (
        request.POST.get('name_ko') or ''
    ).strip()

    name_en = (
        request.POST.get('name_en') or ''
    ).strip()

    if not name_ko:
        messages.error(
            request,
            '한글 견종명을 입력해주세요.'
        )
        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    # 같은 한글명 중복 등록 방지
    if DogBreed.objects.filter(
        name_ko__iexact=name_ko
    ).exists():
        messages.error(
            request,
            f'이미 등록된 견종입니다: {name_ko}'
        )
        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    # slug 생성
    #
    # 한글 slug를 허용하는 모델이므로
    # 견종명을 기본값으로 사용하고 중복 시 숫자를 붙인다.
    from django.utils.text import slugify

    base_slug = slugify(
        name_ko,
        allow_unicode=True
    ) or 'breed'

    slug = base_slug
    seq = 2

    while DogBreed.objects.filter(
        slug=slug
    ).exists():
        slug = f'{base_slug}-{seq}'
        seq += 1

    fci_group_raw = (
        request.POST.get('fci_group') or ''
    ).strip()

    try:
        fci_group = (
            int(fci_group_raw)
            if fci_group_raw
            else None
        )
    except ValueError:
        fci_group = None

    obj = DogBreed(
        name_ko=name_ko,
        name_en=name_en,
        slug=slug,

        fci_group=fci_group,
        fci_group_name=(
            request.POST.get('fci_group_name') or ''
        ).strip(),
        fci_section=(
            request.POST.get('fci_section') or ''
        ).strip(),
        fci_standard_no=(
            request.POST.get('fci_standard_no') or ''
        ).strip(),

        origin=(
            request.POST.get('origin') or ''
        ).strip(),
        patronage=(
            request.POST.get('patronage') or ''
        ).strip(),
        use=(
            request.POST.get('use') or ''
        ).strip(),
        classification=(
            request.POST.get('classification') or ''
        ).strip(),

        summary=(
            request.POST.get('summary') or ''
        ).strip(),
        history=(
            request.POST.get('history') or ''
        ).strip(),
        appearance=(
            request.POST.get('appearance') or ''
        ).strip(),
        temperament=(
            request.POST.get('temperament') or ''
        ).strip(),
        description=(
            request.POST.get('description') or ''
        ).strip(),

        height_male=(
            request.POST.get('height_male') or ''
        ).strip(),
        height_female=(
            request.POST.get('height_female') or ''
        ).strip(),
        purpose_detail=(
            request.POST.get('purpose_detail') or ''
        ).strip(),

        head=(
            request.POST.get('head') or ''
        ).strip(),
        neck=(
            request.POST.get('neck') or ''
        ).strip(),
        body=(
            request.POST.get('body') or ''
        ).strip(),
        tail=(
            request.POST.get('tail') or ''
        ).strip(),
        limbs=(
            request.POST.get('limbs') or ''
        ).strip(),
        gait=(
            request.POST.get('gait') or ''
        ).strip(),
        coat=(
            request.POST.get('coat') or ''
        ).strip(),
        size_detail=(
            request.POST.get('size_detail') or ''
        ).strip(),
        faults=(
            request.POST.get('faults') or ''
        ).strip(),
        disqualification=(
            request.POST.get('disqualification') or ''
        ).strip(),

        image_url=(
            request.POST.get('image_url') or ''
        ).strip(),
        image_source=(
            request.POST.get('image_source') or ''
        ).strip(),

        source_name=(
            request.POST.get('source_name') or ''
        ).strip(),
        source_url=(
            request.POST.get('source_url') or ''
        ).strip(),

        is_published=(
            'is_published' in request.POST
        ),

        sort_order=int(
            request.POST.get('sort_order') or 0
        ),
    )

    if request.FILES.get('image'):
        obj.image = request.FILES['image']

    obj.save()

    messages.success(
        request,
        f'견종 DB에 "{obj.name_ko}"이(가) 등록되었습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


@login_required
@require_POST
def breed_page_create(request):

    title = (request.POST.get('title') or '').strip()

    if not title:
        messages.error(
            request,
            '제목을 입력해주세요.'
        )
        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    obj = ContentItem(
        kind='breed',
        title=title,
        link=(request.POST.get('link') or '').strip(),
        sort_order=int(
            request.POST.get('sort_order') or 0
        ),
        is_published=(
            'is_published' in request.POST
        ),
    )

    if request.FILES.get('image'):
        obj.image = request.FILES['image']

    obj.save()

    messages.success(
        request,
        '전 세계 견종 카드가 등록되었습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


@login_required
@require_POST
def breed_page_update(request, pk):

    obj = get_object_or_404(
        ContentItem,
        pk=pk,
        kind='breed'
    )

    title = (request.POST.get('title') or '').strip()

    if not title:
        messages.error(
            request,
            '제목을 입력해주세요.'
        )
        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    obj.title = title
    obj.link = (
        request.POST.get('link') or ''
    ).strip()

    obj.sort_order = int(
        request.POST.get('sort_order') or 0
    )

    obj.is_published = (
        'is_published' in request.POST
    )

    if request.FILES.get('image'):
        obj.image = request.FILES['image']

    obj.save()

    messages.success(
        request,
        '카드가 수정되었습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


@login_required
@require_POST
def breed_page_bulk_delete(request):

    ids = request.POST.getlist('selected')

    if not ids:
        messages.error(
            request,
            '삭제할 카드를 선택해주세요.'
        )

        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    qs = ContentItem.objects.filter(
        kind='breed',
        pk__in=ids
    )

    count = qs.count()

    qs.delete()

    messages.success(
        request,
        f'{count}개의 카드가 삭제되었습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


# ============================================================
# DOG BREED DATABASE - ADMIN CRUD
# ============================================================

@login_required
@require_POST
def dogbreed_update(request, pk):

    from django.shortcuts import get_object_or_404
    from django.contrib import messages
    from django.urls import reverse
    from django.utils.text import slugify

    obj = get_object_or_404(
        DogBreed,
        pk=pk,
    )

    name_ko = request.POST.get(
        'name_ko',
        ''
    ).strip()

    if not name_ko:
        messages.error(
            request,
            '한글 견종명은 필수입니다.'
        )

        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    # --------------------------------------------------------
    # 기본정보
    # --------------------------------------------------------

    obj.name_ko = name_ko

    obj.name_en = request.POST.get(
        'name_en',
        ''
    ).strip()

    obj.origin = request.POST.get(
        'origin',
        ''
    ).strip()

    obj.patronage = request.POST.get(
        'patronage',
        ''
    ).strip()

    obj.fci_group_name = request.POST.get(
        'fci_group_name',
        ''
    ).strip()

    obj.fci_section = request.POST.get(
        'fci_section',
        ''
    ).strip()

    obj.fci_standard_no = request.POST.get(
        'fci_standard_no',
        ''
    ).strip()

    # --------------------------------------------------------
    # FCI 그룹
    # --------------------------------------------------------

    fci_group = request.POST.get(
        'fci_group',
        ''
    ).strip()

    try:
        obj.fci_group = (
            int(fci_group)
            if fci_group
            else None
        )
    except ValueError:
        obj.fci_group = None

    # --------------------------------------------------------
    # 설명
    # --------------------------------------------------------

    text_fields = [
        'use',
        'classification',
        'summary',
        'history',
        'appearance',
        'temperament',
        'description',
        'purpose_detail',
        'head',
        'neck',
        'body',
        'tail',
        'limbs',
        'gait',
        'coat',
        'size_detail',
        'faults',
        'disqualification',
    ]

    for field in text_fields:

        setattr(
            obj,
            field,
            request.POST.get(
                field,
                ''
            ).strip()
        )

    # --------------------------------------------------------
    # 체고
    # --------------------------------------------------------

    obj.height_male = request.POST.get(
        'height_male',
        ''
    ).strip()

    obj.height_female = request.POST.get(
        'height_female',
        ''
    ).strip()

    # --------------------------------------------------------
    # 이미지
    # --------------------------------------------------------

    uploaded_image = request.FILES.get(
        'image'
    )

    if uploaded_image:
        obj.image = uploaded_image

    obj.image_url = request.POST.get(
        'image_url',
        ''
    ).strip()

    obj.image_source = request.POST.get(
        'image_source',
        ''
    ).strip()

    # --------------------------------------------------------
    # 출처
    # --------------------------------------------------------

    obj.source_name = request.POST.get(
        'source_name',
        ''
    ).strip()

    obj.source_url = request.POST.get(
        'source_url',
        ''
    ).strip()

    obj.source_detail_url = request.POST.get(
        'source_detail_url',
        ''
    ).strip()

    obj.source_id = request.POST.get(
        'source_id',
        ''
    ).strip()

    source_group_id = request.POST.get(
        'source_group_id',
        ''
    ).strip()

    try:
        obj.source_group_id = (
            int(source_group_id)
            if source_group_id
            else None
        )
    except ValueError:
        obj.source_group_id = None

    # --------------------------------------------------------
    # 공개 / 정렬
    # --------------------------------------------------------

    obj.is_published = (
        request.POST.get(
            'is_published'
        )
        == 'on'
    )

    try:
        obj.sort_order = int(
            request.POST.get(
                'sort_order',
                '0'
            )
            or 0
        )
    except ValueError:
        obj.sort_order = 0

    obj.save()

    messages.success(
        request,
        f'{obj.name_ko} 견종 정보를 수정했습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


@login_required
@require_POST
def dogbreed_delete(request, pk):

    from django.shortcuts import get_object_or_404
    from django.contrib import messages

    obj = get_object_or_404(
        DogBreed,
        pk=pk,
    )

    name = obj.name_ko

    obj.delete()

    messages.success(
        request,
        f'{name} 견종을 삭제했습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


@login_required
@require_POST
def dogbreed_bulk_delete(request):

    from django.contrib import messages

    selected = request.POST.getlist(
        'selected_dogbreed'
    )

    ids = []

    for value in selected:

        try:
            ids.append(
                int(value)
            )
        except (TypeError, ValueError):
            pass

    if not ids:

        messages.error(
            request,
            '삭제할 견종을 선택해주세요.'
        )

        return redirect(
            'operator_dashboard:breed_page_editor'
        )

    qs = DogBreed.objects.filter(
        pk__in=ids
    )

    count = qs.count()

    qs.delete()

    messages.success(
        request,
        f'{count}개 견종을 삭제했습니다.'
    )

    return redirect(
        'operator_dashboard:breed_page_editor'
    )


# ============================================================
# BOARD ADMIN CMS
# 공지사항 / 자유게시판 통합 관리
# ============================================================

@login_required
def board_admin_list(request, board_type):

    from django.core.paginator import Paginator
    from django.db.models import Q, Count
    from django.shortcuts import render
    from .models import BoardPost, BoardCategory

    if board_type not in ('notice', 'community'):
        raise Http404

    keyword = request.GET.get('q', '').strip()
    category_slug = request.GET.get('category', '').strip()
    status = request.GET.get('status', '').strip()

    posts = (
        BoardPost.objects
        .filter(board_type=board_type)
        .select_related('category', 'author')
        .annotate(
            comment_count=Count(
                'comments',
                distinct=True
            ),
            attachment_count=Count(
                'attachments',
                distinct=True
            ),
            like_count=Count(
                'likes',
                distinct=True
            ),
        )
        .order_by(
            '-is_pinned',
            '-created_at'
        )
    )

    if keyword:
        posts = posts.filter(
            Q(title__icontains=keyword)
            | Q(body__icontains=keyword)
        )

    if category_slug:
        posts = posts.filter(
            category__slug=category_slug
        )

    if status == 'published':
        posts = posts.filter(
            is_published=True,
            is_hidden=False
        )

    elif status == 'hidden':
        posts = posts.filter(
            is_hidden=True
        )

    elif status == 'draft':
        posts = posts.filter(
            is_published=False
        )

    categories = (
        BoardCategory.objects
        .filter(board_type=board_type)
        .order_by('sort_order', 'id')
    )

    paginator = Paginator(posts, 20)

    page_obj = paginator.get_page(
        request.GET.get('page')
    )

    context = {
        'board_type': board_type,
        'board_label': (
            '공지사항'
            if board_type == 'notice'
            else '자유게시판'
        ),
        'page_obj': page_obj,
        'categories': categories,
        'keyword': keyword,
        'selected_category': category_slug,
        'selected_status': status,
        'total_count': posts.count(),
    }

    return render(
        request,
        'dashboard/board_admin_list.html',
        context
    )


@login_required
def board_admin_create(request, board_type):

    from django.contrib import messages
    from django.shortcuts import render, redirect
    from django.utils import timezone
    from .models import BoardPost, BoardCategory

    if board_type not in ('notice', 'community'):
        raise Http404

    categories = (
        BoardCategory.objects
        .filter(
            board_type=board_type,
            is_active=True
        )
        .order_by('sort_order', 'id')
    )

    board_label = (
        '공지사항'
        if board_type == 'notice'
        else '자유게시판'
    )

    if request.method == 'POST':

        title = request.POST.get(
            'title',
            ''
        ).strip()

        body = request.POST.get(
            'body',
            ''
        ).strip()

        category_id = request.POST.get(
            'category',
            ''
        ).strip()

        if not title:

            messages.error(
                request,
                '제목을 입력해주세요.'
            )

            return render(
                request,
                'dashboard/board_admin_form.html',
                {
                    'board_type': board_type,
                    'board_label': board_label,
                    'categories': categories,
                    'mode': 'create',
                }
            )

        category = None

        if category_id:

            category = (
                BoardCategory.objects
                .filter(
                    pk=category_id,
                    board_type=board_type
                )
                .first()
            )

        is_published = (
            request.POST.get('is_published')
            == 'on'
        )

        post = BoardPost.objects.create(
            board_type=board_type,
            category=category,
            author=request.user,
            title=title,
            body=body,

            is_published=is_published,

            is_pinned=(
                request.POST.get('is_pinned')
                == 'on'
            ),

            is_important=(
                request.POST.get('is_important')
                == 'on'
            ),

            is_featured=(
                request.POST.get('is_featured')
                == 'on'
            ),

            allow_comments=(
                request.POST.get('allow_comments')
                == 'on'
            ),

            published_at=(
                timezone.now()
                if is_published
                else None
            ),
        )

        # 첨부파일 저장
        _save_board_attachments(
            request,
            post
        )

        messages.success(
            request,
            '게시글을 등록했습니다.'
        )

        return redirect(
            'operator_dashboard:board_admin_list',
            board_type=board_type
        )

    return render(
        request,
        'dashboard/board_admin_form.html',
        {
            'board_type': board_type,
            'board_label': board_label,
            'categories': categories,
            'mode': 'create',
        }
    )


@login_required
def board_admin_update(request, board_type, pk):

    from django.contrib import messages
    from django.shortcuts import (
        render,
        redirect,
        get_object_or_404,
    )
    from django.utils import timezone
    from .models import (
        BoardPost,
        BoardCategory,
        BoardAttachment,
    )

    if board_type not in ('notice', 'community'):
        raise Http404

    post = get_object_or_404(
        BoardPost,
        pk=pk,
        board_type=board_type
    )

    categories = (
        BoardCategory.objects
        .filter(board_type=board_type)
        .order_by('sort_order', 'id')
    )

    if request.method == 'POST':

        title = request.POST.get(
            'title',
            ''
        ).strip()

        body = request.POST.get(
            'body',
            ''
        ).strip()

        if not title:

            messages.error(
                request,
                '제목을 입력해주세요.'
            )

            return redirect(
                'operator_dashboard:board_admin_update',
                board_type=board_type,
                pk=post.pk
            )

        category_id = request.POST.get(
            'category',
            ''
        ).strip()

        category = None

        if category_id:
            category = (
                BoardCategory.objects
                .filter(
                    pk=category_id,
                    board_type=board_type
                )
                .first()
            )

        old_published = post.is_published

        post.category = category
        post.title = title
        post.body = body

        post.is_published = (
            request.POST.get('is_published')
            == 'on'
        )

        post.is_pinned = (
            request.POST.get('is_pinned')
            == 'on'
        )

        post.is_important = (
            request.POST.get('is_important')
            == 'on'
        )

        post.is_featured = (
            request.POST.get('is_featured')
            == 'on'
        )

        post.allow_comments = (
            request.POST.get('allow_comments')
            == 'on'
        )

        if (
            post.is_published
            and not old_published
        ):
            post.published_at = timezone.now()

        post.save()

        delete_attachment_ids = request.POST.getlist(
            'delete_attachment'
        )

        if delete_attachment_ids:

            post.attachments.filter(
                pk__in=delete_attachment_ids
            ).delete()

        files = request.FILES.getlist(
            'attachments'
        )

        import os

        image_extensions = {
            '.jpg',
            '.jpeg',
            '.png',
            '.gif',
            '.webp',
            '.bmp',
        }

        current_count = post.attachments.count()

        for index, uploaded in enumerate(files):

            ext = os.path.splitext(
                uploaded.name
            )[1].lower()

            BoardAttachment.objects.create(
                post=post,
                file=uploaded,
                original_name=uploaded.name,
                is_image=(
                    ext in image_extensions
                ),
                sort_order=current_count + index
            )

        messages.success(
            request,
            '게시글을 수정했습니다.'
        )

        _save_board_attachments(request, post)

    return redirect(
            'operator_dashboard:board_admin_list',
            board_type=board_type
        )

    return render(
        request,
        'dashboard/board_admin_form.html',
        {
            'board_type': board_type,
            'board_label': (
                '공지사항'
                if board_type == 'notice'
                else '자유게시판'
            ),
            'categories': categories,
            'post': post,
            'mode': 'update',
        }
    )


@login_required
@require_POST
def board_admin_delete(request, board_type, pk):

    from django.contrib import messages
    from django.shortcuts import (
        redirect,
        get_object_or_404,
    )
    from .models import BoardPost

    if board_type not in ('notice', 'community'):
        raise Http404

    post = get_object_or_404(
        BoardPost,
        pk=pk,
        board_type=board_type
    )

    title = post.title

    post.delete()

    messages.success(
        request,
        f'{title} 게시글을 삭제했습니다.'
    )

    return redirect(
        'operator_dashboard:board_admin_list',
        board_type=board_type
    )


@login_required
@require_POST
def board_admin_bulk_action(request, board_type):

    from django.contrib import messages
    from django.shortcuts import redirect
    from .models import BoardPost

    if board_type not in ('notice', 'community'):
        raise Http404

    selected = request.POST.getlist(
        'selected_posts'
    )

    ids = []

    for value in selected:

        try:
            ids.append(int(value))

        except (TypeError, ValueError):
            pass

    if not ids:

        messages.error(
            request,
            '게시글을 선택해주세요.'
        )

        return redirect(
            'operator_dashboard:board_admin_list',
            board_type=board_type
        )

    action = request.POST.get(
        'bulk_action',
        ''
    ).strip()

    qs = BoardPost.objects.filter(
        board_type=board_type,
        pk__in=ids
    )

    count = qs.count()

    if action == 'publish':

        qs.update(
            is_published=True,
            is_hidden=False
        )

        message = f'{count}개 게시글을 공개했습니다.'

    elif action == 'hide':

        qs.update(
            is_hidden=True
        )

        message = f'{count}개 게시글을 숨겼습니다.'

    elif action == 'pin':

        qs.update(
            is_pinned=True
        )

        message = f'{count}개 게시글을 상단 고정했습니다.'

    elif action == 'unpin':

        qs.update(
            is_pinned=False
        )

        message = f'{count}개 게시글의 상단 고정을 해제했습니다.'

    elif action == 'delete':

        qs.delete()

        message = f'{count}개 게시글을 삭제했습니다.'

    else:

        messages.error(
            request,
            '올바른 관리 작업을 선택해주세요.'
        )

        return redirect(
            'operator_dashboard:board_admin_list',
            board_type=board_type
        )

    messages.success(
        request,
        message
    )

    return redirect(
        'operator_dashboard:board_admin_list',
        board_type=board_type
    )


# ============================================================
# BOARD CATEGORY ADMIN
# ============================================================

@login_required
def board_category_admin(request, board_type):

    from django.contrib import messages
    from django.shortcuts import render, redirect
    from django.utils.text import slugify
    from .models import BoardCategory

    if board_type not in ('notice', 'community'):
        raise Http404

    if request.method == 'POST':

        name = request.POST.get(
            'name',
            ''
        ).strip()

        slug = request.POST.get(
            'slug',
            ''
        ).strip()

        sort_order = request.POST.get(
            'sort_order',
            '0'
        ).strip()

        if not name:

            messages.error(
                request,
                '카테고리명을 입력해주세요.'
            )

            return redirect(
                'operator_dashboard:board_category_admin',
                board_type=board_type
            )

        if not slug:

            slug = slugify(
                name,
                allow_unicode=True
            )

        base_slug = slug
        suffix = 2

        while BoardCategory.objects.filter(
            board_type=board_type,
            slug=slug
        ).exists():

            slug = f'{base_slug}-{suffix}'
            suffix += 1

        try:
            sort_order = int(sort_order)

        except (TypeError, ValueError):
            sort_order = 0

        BoardCategory.objects.create(
            board_type=board_type,
            name=name,
            slug=slug,
            sort_order=sort_order,
            is_active=True
        )

        messages.success(
            request,
            f'{name} 카테고리를 추가했습니다.'
        )

        return redirect(
            'operator_dashboard:board_category_admin',
            board_type=board_type
        )

    categories = (
        BoardCategory.objects
        .filter(board_type=board_type)
        .annotate(
            post_count=Count('posts')
        )
        .order_by(
            'sort_order',
            'id'
        )
    )

    return render(
        request,
        'dashboard/board_category_admin.html',
        {
            'board_type': board_type,
            'board_label': (
                '공지사항'
                if board_type == 'notice'
                else '자유게시판'
            ),
            'categories': categories,
        }
    )


@login_required
@require_POST
def board_category_update(request, board_type, pk):

    from django.contrib import messages
    from django.shortcuts import (
        redirect,
        get_object_or_404,
    )
    from .models import BoardCategory

    if board_type not in ('notice', 'community'):
        raise Http404

    category = get_object_or_404(
        BoardCategory,
        pk=pk,
        board_type=board_type
    )

    name = request.POST.get(
        'name',
        ''
    ).strip()

    if name:
        category.name = name

    try:
        category.sort_order = int(
            request.POST.get(
                'sort_order',
                category.sort_order
            )
        )

    except (TypeError, ValueError):
        pass

    category.is_active = (
        request.POST.get('is_active')
        == 'on'
    )

    category.save()

    messages.success(
        request,
        '카테고리를 수정했습니다.'
    )

    return redirect(
        'operator_dashboard:board_category_admin',
        board_type=board_type
    )


@login_required
@require_POST
def board_category_delete(request, board_type, pk):

    from django.contrib import messages
    from django.shortcuts import (
        redirect,
        get_object_or_404,
    )
    from .models import BoardCategory

    if board_type not in ('notice', 'community'):
        raise Http404

    category = get_object_or_404(
        BoardCategory,
        pk=pk,
        board_type=board_type
    )

    name = category.name

    category.delete()

    messages.success(
        request,
        f'{name} 카테고리를 삭제했습니다.'
    )

    return redirect(
        'operator_dashboard:board_category_admin',
        board_type=board_type
    )


# ============================================================
# BOARD ATTACHMENT COMMON HELPER
# ============================================================

def _save_board_attachments(request, post):

    from .models import BoardAttachment

    files = []

    for key in request.FILES.keys():

        for uploaded in request.FILES.getlist(key):

            if uploaded:
                files.append(uploaded)

    if not files:
        return 0

    field_names = {
        f.name
        for f in BoardAttachment._meta.fields
    }

    count = 0

    for uploaded in files:

        values = {}

        if 'post' in field_names:
            values['post'] = post

        if 'file' in field_names:
            values['file'] = uploaded
        else:
            continue

        if 'original_name' in field_names:
            values['original_name'] = uploaded.name

        if 'file_name' in field_names:
            values['file_name'] = uploaded.name

        if 'file_size' in field_names:
            values['file_size'] = uploaded.size

        if 'content_type' in field_names:
            values['content_type'] = (
                uploaded.content_type or ''
            )

        BoardAttachment.objects.create(
            **values
        )

        count += 1

    return count


# ============================================================
# BOARD REPORT ADMIN
# ============================================================

@login_required
def board_report_admin_list(request):

    from django.core.paginator import Paginator
    from django.shortcuts import render
    from django.db.models import Q
    from .models import BoardReport

    keyword = request.GET.get(
        'q',
        ''
    ).strip()

    selected_status = request.GET.get(
        'status',
        ''
    ).strip()

    reports = (
        BoardReport.objects
        .select_related(
            'post',
            'post__author',
            'reporter',
        )
        .order_by('-created_at', '-pk')
    )

    if keyword:
        reports = reports.filter(
            Q(post__title__icontains=keyword)
            | Q(reason__icontains=keyword)
            | Q(detail__icontains=keyword)
            | Q(reporter__username__icontains=keyword)
        )

    if selected_status:
        reports = reports.filter(
            status=selected_status
        )

    status_field = BoardReport._meta.get_field(
        'status'
    )

    status_choices = list(
        status_field.choices or []
    )

    paginator = Paginator(
        reports,
        20
    )

    page_obj = paginator.get_page(
        request.GET.get('page')
    )

    return render(
        request,
        'dashboard/board_report_admin.html',
        {
            'page_obj': page_obj,
            'total_count': reports.count(),
            'keyword': keyword,
            'selected_status': selected_status,
            'status_choices': status_choices,
        }
    )


@login_required
def board_report_admin_detail(request, pk):

    from django.contrib import messages
    from django.shortcuts import (
        render,
        redirect,
        get_object_or_404,
    )
    from .models import (
        BoardReport,
        BoardPost,
    )

    report = get_object_or_404(
        BoardReport.objects.select_related(
            'post',
            'post__author',
            'reporter',
        ),
        pk=pk
    )

    status_field = BoardReport._meta.get_field(
        'status'
    )

    status_choices = list(
        status_field.choices or []
    )

    if request.method == 'POST':

        action = request.POST.get(
            'action',
            ''
        ).strip()

        if action == 'status':

            new_status = request.POST.get(
                'status',
                ''
            ).strip()

            valid_statuses = {
                value
                for value, label
                in status_choices
            }

            # choices가 없는 모델도 기존 CharField 값은
            # 빈 값이 아닌 경우 처리 가능하게 둔다.
            if (
                new_status
                and (
                    not valid_statuses
                    or new_status in valid_statuses
                )
            ):
                report.status = new_status
                report.save(
                    update_fields=['status']
                )

                messages.success(
                    request,
                    '신고 처리 상태를 변경했습니다.'
                )

            else:
                messages.error(
                    request,
                    '올바르지 않은 처리 상태입니다.'
                )

        elif action == 'hide_post':

            BoardPost.objects.filter(
                pk=report.post_id
            ).update(
                is_hidden=True
            )

            messages.success(
                request,
                '신고된 게시글을 숨김 처리했습니다.'
            )

        elif action == 'restore_post':

            BoardPost.objects.filter(
                pk=report.post_id
            ).update(
                is_hidden=False
            )

            messages.success(
                request,
                '게시글 숨김을 해제했습니다.'
            )

        else:

            messages.error(
                request,
                '처리할 작업을 선택해주세요.'
            )

        return redirect(
            'operator_dashboard:board_report_admin_detail',
            pk=report.pk
        )

    return render(
        request,
        'dashboard/board_report_admin_detail.html',
        {
            'report': report,
            'status_choices': status_choices,
        }
    )
