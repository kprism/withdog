from io import BytesIO
from django.contrib.admin.views.decorators import staff_member_required
from django.http import JsonResponse
from django.views.decorators.http import require_GET,require_POST
from django.utils import timezone
from django.db.models import Q
from openpyxl import load_workbook
from .models import Partner,PartnerCategory,ContentItem,SiteSetting,IntroLayer
@require_GET
def partner_list(request):
 qs=Partner.objects.filter(is_active=True).select_related('category'); category=request.GET.get('category'); city=request.GET.get('city'); q=request.GET.get('q')
 if category:qs=qs.filter(category__code=category)
 if city:qs=qs.filter(city=city)
 if q:qs=qs.filter(name__icontains=q)
 data=[{'id':p.id,'name':p.name,'address':p.address,'phone':p.phone,'city':p.city,'category':p.category.code,'category_name':p.category.name,'latitude':float(p.latitude) if p.latitude is not None else None,'longitude':float(p.longitude) if p.longitude is not None else None,'is_affiliated':p.is_affiliated} for p in qs[:2000]]
 return JsonResponse({'count':qs.count(),'results':data})
@require_GET
def content_feed(request):
 kind=request.GET.get('kind')
 qs=ContentItem.objects.filter(is_published=True)
 if kind:qs=qs.filter(kind=kind)
 if kind == 'hero':qs=qs.prefetch_related('intro_layers')
 now=timezone.now()
 normal=['notice','community','breed','intro','hero']
 qs=qs.filter(Q(kind__in=normal)|Q(kind='popup',popup_start__isnull=True)|Q(kind='popup',popup_start__lte=now)).filter(Q(kind__in=normal)|Q(kind='popup',popup_end__isnull=True)|Q(kind='popup',popup_end__gte=now))
 setting=SiteSetting.get_solo()
 results=[]
 for x in qs[:100]:
  layers=[]
  if x.kind == 'hero':
   for layer in x.intro_layers.all():
    if not layer.is_visible:continue
    layers.append({
     'id':layer.id,'name':layer.name,'layer_type':layer.layer_type,'text':layer.text,'link':layer.link,
     'image':layer.image.url if layer.image else '','video':layer.video.url if layer.video else '',
     'x_px':layer.x_px,'y_px':layer.y_px,'width_px':layer.width_px,'height_px':layer.height_px,
     'opacity':layer.opacity,'z_index':layer.z_index,'font_size':layer.font_size,'font_weight':layer.font_weight,
     'color':layer.color,'background':layer.background,'border_radius':layer.border_radius,'object_fit':layer.object_fit,
     'animation':layer.animation,'animation_delay_ms':layer.animation_delay_ms,'animation_duration_ms':layer.animation_duration_ms,
    })
  results.append({
   'id':x.id,'kind':x.kind,'title':x.title,'label':x.label,'body':x.body,
   'image':x.image.url if x.image else x.image_url,'link':x.link,'media_type':x.media_type,
   'video':x.video.url if x.video else '','video_url':x.video_url,'autoplay':x.autoplay,'muted':x.muted,'loop':x.loop,
   'sort_order':x.sort_order,'hero_eyebrow':x.hero_eyebrow,'hero_title':x.hero_title,'hero_meta1':x.hero_meta1,
   'hero_meta2':x.hero_meta2,'hero_meta1_font_size':x.hero_meta1_font_size,'hero_meta2_font_size':x.hero_meta2_font_size,
   'hero_button1_text':x.hero_button1_text,'hero_button1_link':x.hero_button1_link,
   'hero_button2_text':x.hero_button2_text,'hero_button2_link':x.hero_button2_link,
   'stage_width':setting.intro_stage_width,'stage_height':setting.intro_stage_height,'layers':layers,
   'created_at':x.created_at.strftime('%Y-%m-%d'),
  })
 return JsonResponse({'count':len(results),'results':results})

@staff_member_required
@require_POST
def import_excel(request):
 f=request.FILES.get('file');
 if not f:return JsonResponse({'error':'엑셀 파일이 필요합니다.'},status=400)
 cat,_=PartnerCategory.objects.get_or_create(code=request.POST.get('category','grooming'),defaults={'name':request.POST.get('category_name','미용센터')}); wb=load_workbook(BytesIO(f.read()),data_only=True); created=updated=skipped=0
 for ws in wb.worksheets:
  city=ws.title.rstrip('0123456789')
  for row in ws.iter_rows(min_row=4,values_only=True):
   if len(row)<3 or not row[1] or not row[2]:skipped+=1;continue
   _,new=Partner.objects.update_or_create(category=cat,name=str(row[1]).strip(),address=str(row[2]).strip(),defaults={'phone':str(row[3]).strip() if len(row)>3 and row[3] else '','city':city,'source':f.name,'is_active':True});created+=int(new);updated+=int(not new)
 return JsonResponse({'ok':True,'created':created,'updated':updated,'skipped':skipped})
@staff_member_required
@require_POST
def save_coordinates(request,pk):
 try:p=Partner.objects.get(pk=pk)
 except Partner.DoesNotExist:return JsonResponse({'error':'not found'},status=404)
 p.latitude=request.POST.get('latitude') or None;p.longitude=request.POST.get('longitude') or None;p.save(update_fields=['latitude','longitude','updated_at']);return JsonResponse({'ok':True})


# === PUBLIC MEMBER REGISTER API ===
import json
from datetime import datetime

from django.contrib.auth import get_user_model
from django.db import transaction
from django.http import JsonResponse
from django.views.decorators.http import require_POST

from .models import MemberProfile


@require_POST
def member_register(request):
    try:
        data = json.loads(request.body.decode("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse(
            {"ok": False, "message": "잘못된 요청입니다."},
            status=400,
        )

    name = str(data.get("name", "")).strip()
    birth_date = str(data.get("birth_date", "")).strip()
    gender = str(data.get("gender", "")).strip()
    phone = str(data.get("phone", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))
    region = str(data.get("region", "")).strip()
    address_detail = str(data.get("address_detail", "")).strip()

    privacy_agreed = bool(data.get("privacy_agreed"))
    terms_agreed = bool(data.get("terms_agreed"))
    rules_agreed = bool(data.get("rules_agreed"))
    regular_member_requested = bool(data.get("regular_member_requested"))

    if not name:
        return JsonResponse(
            {"ok": False, "message": "이름을 입력해 주세요."},
            status=400,
        )

    if not phone:
        return JsonResponse(
            {"ok": False, "message": "전화번호를 입력해 주세요."},
            status=400,
        )

    if not email:
        return JsonResponse(
            {"ok": False, "message": "이메일을 입력해 주세요."},
            status=400,
        )

    if not password or len(password) < 6:
        return JsonResponse(
            {"ok": False, "message": "비밀번호는 6자 이상 입력해 주세요."},
            status=400,
        )

    if not (privacy_agreed and terms_agreed and rules_agreed):
        return JsonResponse(
            {"ok": False, "message": "필수 방침·약관·회원규칙 동의가 필요합니다."},
            status=400,
        )

    if gender not in ("M", "F", "O", ""):
        return JsonResponse(
            {"ok": False, "message": "성별 값이 올바르지 않습니다."},
            status=400,
        )

    parsed_birth = None
    if birth_date:
        for fmt in ("%Y-%m-%d", "%Y.%m.%d"):
            try:
                parsed_birth = datetime.strptime(
                    birth_date, fmt
                ).date()
                break
            except ValueError:
                pass

        if parsed_birth is None:
            return JsonResponse(
                {
                    "ok": False,
                    "message": "생년월일 형식이 올바르지 않습니다.",
                },
                status=400,
            )

    User = get_user_model()

    if User.objects.filter(email__iexact=email).exists():
        return JsonResponse(
            {
                "ok": False,
                "message": "이미 가입된 이메일입니다.",
            },
            status=409,
        )

    if User.objects.filter(username__iexact=email).exists():
        return JsonResponse(
            {
                "ok": False,
                "message": "이미 가입된 이메일입니다.",
            },
            status=409,
        )

    try:
        with transaction.atomic():
            user = User.objects.create_user(
                username=email,
                email=email,
                password=password,
            )

            MemberProfile.objects.create(
                user=user,
                name=name,
                birth_date=parsed_birth,
                gender=gender,
                phone=phone,
                region=region,
                address_detail=address_detail,
                privacy_agreed=privacy_agreed,
                terms_agreed=terms_agreed,
                rules_agreed=rules_agreed,
                agreements_updated_at=timezone.now(),
                regular_member_requested=regular_member_requested,
                membership_status=(
                    "pending"
                    if regular_member_requested
                    else "general"
                ),
            )

    except Exception:
        return JsonResponse(
            {
                "ok": False,
                "message": "회원가입 처리 중 오류가 발생했습니다.",
            },
            status=500,
        )

    result={"ok":True,"message":"가입을 축하드립니다.","email":email,"payment_required":False}
    if regular_member_requested:
        from .payment_views import create_membership_payment
        payment,error=create_membership_payment(user)
        if payment:
            result["payment_required"]=True
            result["payment"]=payment
        else:
            result["payment_message"]=error
    return JsonResponse(result,status=201)


# === PUBLIC CSRF COOKIE API ===
from django.views.decorators.csrf import ensure_csrf_cookie


@ensure_csrf_cookie
def csrf_cookie(request):
    return JsonResponse({
        "ok": True,
        "message": "CSRF cookie ready",
    })


# === PUBLIC LOGIN API ===
from django.contrib.auth import authenticate, login, logout
from django.views.decorators.http import require_GET


@require_POST
def member_login(request):
    try:
        data = json.loads(request.body.decode("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse(
            {"ok": False, "message": "잘못된 요청입니다."},
            status=400,
        )

    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))

    if not email or not password:
        return JsonResponse(
            {"ok": False, "message": "이메일과 비밀번호를 입력해 주세요."},
            status=400,
        )

    user = authenticate(
        request,
        username=email,
        password=password,
    )

    if user is None:
        return JsonResponse(
            {"ok": False, "message": "이메일 또는 비밀번호가 올바르지 않습니다."},
            status=401,
        )

    if not user.is_active:
        return JsonResponse(
            {"ok": False, "message": "탈퇴 또는 비활성 처리된 회원입니다."},
            status=403,
        )

    if user.is_staff:
        return JsonResponse(
            {"ok": False, "message": "관리자 계정은 운영자 로그인 화면을 이용해 주세요."},
            status=403,
        )

    login(request, user)

    profile = getattr(user, "member_profile", None)

    return JsonResponse({
        "ok": True,
        "message": "로그인되었습니다.",
        "user": {
            "email": user.email or user.username,
            "name": profile.name if profile else user.get_full_name(),
            "membership_status": (
                profile.membership_status
                if profile else "general"
            ),
        },
    })


@require_POST
def member_logout(request):
    logout(request)
    return JsonResponse({
        "ok": True,
        "message": "로그아웃되었습니다.",
    })


@require_GET
def member_session(request):
    if not request.user.is_authenticated or request.user.is_staff:
        return JsonResponse({
            "ok": True,
            "authenticated": False,
        })

    profile = getattr(request.user, "member_profile", None)

    return JsonResponse({
        "ok": True,
        "authenticated": True,
        "user": {
            "email": request.user.email or request.user.username,
            "name": profile.name if profile else request.user.get_full_name(),
            "membership_status": (
                profile.membership_status
                if profile else "general"
            ),
        },
    })


# === PUBLIC SITE SETTINGS API ===
@require_GET
def site_settings_api(request):
    x = SiteSetting.get_solo()

    return JsonResponse({
        'ok': True,
        'site_name': x.site_name,
        'site_subtitle': x.site_subtitle,
        'logo': x.logo.url if x.logo else '',
        'header_logo_size': x.header_logo_size,
        'header_site_name_size': x.header_site_name_size,
        'header_site_name_weight': x.header_site_name_weight,
        'header_subtitle_size': x.header_subtitle_size,
        'header_subtitle_weight': x.header_subtitle_weight,
        'header_menu_font_size': x.header_menu_font_size,
        'header_menu_font_weight': x.header_menu_font_weight,
        'intro_header_opacity': x.intro_header_opacity,
        'intro_stage_width': x.intro_stage_width,
        'intro_stage_height': x.intro_stage_height,
        'hero_eyebrow': x.hero_eyebrow,
        'hero_title': x.hero_title,
        'hero_meta1': x.hero_meta1,
        'hero_meta2': x.hero_meta2,
        'hero_button1_text': x.hero_button1_text,
        'hero_button1_link': x.hero_button1_link,
        'hero_button2_text': x.hero_button2_text,
        'hero_button2_link': x.hero_button2_link,
        'updated_at': x.updated_at.isoformat(),
    })


# === MEMBER ACCOUNT RECOVERY API ===
# Append this block to backend/partners/views.py

import re
from django.contrib.auth import get_user_model
from django.views.decorators.http import require_POST
from django.http import JsonResponse

def _recovery_json(request):
    try:
        return json.loads(request.body.decode("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return None

def _norm_phone(value):
    return re.sub(r"\D", "", str(value or ""))

def _find_member_profile(name, phone):
    name = str(name or "").strip()
    phone = _norm_phone(phone)
    if not name or not phone:
        return None

    # 전화번호 저장 형식(하이픈 유무)이 다를 수 있어 이름 후보를 좁힌 뒤 정규화 비교
    for profile in MemberProfile.objects.select_related("user").filter(name=name):
        if _norm_phone(profile.phone) == phone and profile.user.is_active:
            return profile
    return None

@require_POST
def member_find_id(request):
    data = _recovery_json(request)
    if data is None:
        return JsonResponse({"ok": False, "message": "잘못된 요청입니다."}, status=400)

    profile = _find_member_profile(data.get("name"), data.get("phone"))
    if not profile:
        return JsonResponse(
            {"ok": False, "message": "이름과 연락처가 일치하는 회원을 찾을 수 없습니다."},
            status=404,
        )

    return JsonResponse({
        "ok": True,
        "message": "가입 아이디를 확인했습니다.",
        "email": profile.user.email or profile.user.username,
    })

@require_POST
def member_reset_password(request):
    data = _recovery_json(request)
    if data is None:
        return JsonResponse({"ok": False, "message": "잘못된 요청입니다."}, status=400)

    profile = _find_member_profile(data.get("name"), data.get("phone"))
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))

    if not profile or not email:
        return JsonResponse(
            {"ok": False, "message": "회원 정보를 확인할 수 없습니다."},
            status=404,
        )

    user_email = (profile.user.email or profile.user.username or "").strip().lower()
    if user_email != email:
        return JsonResponse(
            {"ok": False, "message": "이름, 연락처, 이메일이 일치하지 않습니다."},
            status=404,
        )

    if len(password) < 6:
        return JsonResponse(
            {"ok": False, "message": "새 비밀번호는 6자 이상 입력해 주세요."},
            status=400,
        )

    profile.user.set_password(password)
    profile.user.save(update_fields=["password"])

    return JsonResponse({
        "ok": True,
        "message": "비밀번호가 변경되었습니다. 새 비밀번호로 로그인해 주세요.",
    })


# ============================================================
# PUBLIC PARTNER MAP API
# ============================================================

def public_partner_map_api(request):
    from django.http import JsonResponse
    from django.db.models import Q
    from django.utils import timezone

    today = timezone.localdate()

    qs = (
        Partner.objects
        .select_related('category')
        .filter(is_active=True)
    )

    category = request.GET.get('category', '').strip()
    city = request.GET.get('city', '').strip()
    q = request.GET.get('q', '').strip()

    if category:
        qs = qs.filter(category__code=category)

    if city:
        qs = qs.filter(city=city)

    if q:
        qs = qs.filter(
            Q(name__icontains=q) |
            Q(address__icontains=q) |
            Q(phone__icontains=q)
        )

    rows = []

    for partner in qs:
        ad_active = False

        if partner.is_affiliated:
            start_ok = (
                partner.ad_start_date is None
                or partner.ad_start_date <= today
            )

            end_ok = (
                partner.ad_end_date is None
                or partner.ad_end_date >= today
            )

            ad_active = start_ok and end_ok

        rows.append({
            'id': partner.pk,

            'name': partner.name,
            'address': partner.address,
            'phone': partner.phone or '',
            'city': partner.city or '',

            'kakao_place_id': partner.kakao_place_id or '',
            'kakao_place_url': partner.kakao_place_url or '',

            'category': {
                'id': partner.category_id,
                'code': partner.category.code,
                'name': partner.category.name,
            },

            'latitude': (
                float(partner.latitude)
                if partner.latitude is not None
                else None
            ),

            'longitude': (
                float(partner.longitude)
                if partner.longitude is not None
                else None
            ),

            'is_ad': ad_active,

            'ad_monthly_fee': (
                partner.ad_monthly_fee
                if ad_active
                else 0
            ),
        })

    # 광고업체 우선
    # 같은 광고업체끼리는 월 광고금액이 높은 업체 우선
    # 이후 지역/상호명 순
    rows.sort(
        key=lambda x: (
            0 if x['is_ad'] else 1,
            -x['ad_monthly_fee'],
            x['city'],
            x['name'],
        )
    )

    categories = list(
        PartnerCategory.objects
        .order_by('sort_order', 'name')
        .values(
            'id',
            'code',
            'name',
        )
    )

    cities = list(
        Partner.objects
        .filter(is_active=True)
        .exclude(city='')
        .values_list('city', flat=True)
        .distinct()
        .order_by('city')
    )

    return JsonResponse({
        'ok': True,
        'count': len(rows),
        'partners': rows,
        'categories': categories,
        'cities': cities,
    })


def _public_member_payload(user):
    p=user.member_profile
    address_short=' '.join((p.region or p.address_detail or '').split()[:2])
    return {'email':user.email or user.username,'name':p.name,'birth_date':p.birth_date.strftime('%Y.%m.%d') if p.birth_date else '','gender':p.gender if p.gender in ('M','F') else '','phone':p.phone,'address_short':address_short,'membership_status':p.membership_status,'member_no':f'GN-{user.pk:06d}'}

@require_GET
def member_mypage(request):
    if not request.user.is_authenticated or request.user.is_staff:return JsonResponse({'ok':False,'message':'로그인이 필요합니다.'},status=401)
    return JsonResponse({'ok':True,'user':_public_member_payload(request.user)})

@require_POST
def member_self_withdraw(request):
    if not request.user.is_authenticated or request.user.is_staff:return JsonResponse({'ok':False,'message':'로그인이 필요합니다.'},status=401)
    user=request.user;p=user.member_profile
    p.membership_status='withdrawn';p.regular_member_requested=False;p.save(update_fields=['membership_status','regular_member_requested'])
    user.is_active=False;user.save(update_fields=['is_active']);logout(request)
    return JsonResponse({'ok':True,'message':'회원 탈퇴가 완료되었습니다.'})
